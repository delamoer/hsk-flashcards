#!/usr/bin/env python3
"""Convert the 课文挖空练习 xlsx files into per-unit JSON for the 课文 section.

Run with uv (openpyxl provided on the fly):
    uv run --with openpyxl python3 scripts/convert_texts.py

Sources: sources/HSK{1,2,3,4}_课文挖空练习.xlsx + sources/新HSK3.0_第{一,二,三}册_课文挖空练习.xlsx
         + sources/新HSK3.0_第四册{上,下}_课文挖空练习.xlsx (two books, one unit — merged)
         + sources/HSK{5,6}_课文原文.xlsx (essays, no cloze — 读原文 + 连词成句 only)
         + sources/走遍中国_第四册_课文原文.xlsx (dialogue + readings per 单元, gloss in the xlsx)
  新版 → series "newhsk3", 旧版 → series "hsk" (same taxonomy as the flashcard courses).

Output (kept in a subfolder so the flashcard glob `./*-*.json` never picks it up):
  src/data/texts/{series}-{unit}.json
  src/data/texts/meta.json     {"{series}-{unit}": {lessonCount, textCount}}

Each text carries TWO cloze layers ("双重挖空"):
  vocab   — word-level: numbered blanks （1）____, filled from a word bank
  grammar — sentence-level: one blank spanning a whole clause, rebuilt by
            arranging shuffled word tiles (连词成句). The word segmentation is
            NOT auto-generated: it's a hand-vetted table in
            scripts/grammar_segments.tsv (edit that, not this script).
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "data" / "texts"
DATA = ROOT / "src" / "data"
SEG_FILE = ROOT / "scripts" / "grammar_segments.tsv"
MERGE_FILE = ROOT / "scripts" / "merge_words.txt"
GLOSS_FILE = ROOT / "scripts" / "text_gloss.tsv"
WORD_GLOSS_FILE = ROOT / "scripts" / "word_gloss.tsv"
MAX_WORD = 6  # longest word to try in max-match segmentation
MAX_SENT_TOKENS = 15  # 连词成句: sentences longer than this are dropped (too hard)

# (filename, series, unit) — teacher sheet "教师答案版" is the complete one.
SOURCES = [
    ("HSK1_课文挖空练习.xlsx", "hsk", 1),
    ("HSK2_课文挖空练习.xlsx", "hsk", 2),
    ("HSK3_课文挖空练习.xlsx", "hsk", 3),
    ("新HSK3.0_第一册_课文挖空练习.xlsx", "newhsk3", 1),
    ("新HSK3.0_第二册_课文挖空练习.xlsx", "newhsk3", 2),
    ("新HSK3.0_第三册_课文挖空练习.xlsx", "newhsk3", 3),
    # 第四册 = 上(1–10课) + 下(11–20课) 两本书合成一个 unit，main() 按课次合并。
    # 两本都无拼音/英文列 → 每行 py/en 为空，ReadPanel 的 拼音/EN pill 自动隐藏。
    ("新HSK3.0_第四册上_课文挖空练习.xlsx", "newhsk3", 4),
    ("新HSK3.0_第四册下_课文挖空练习.xlsx", "newhsk3", 4),
    ("HSK4_课文挖空练习.xlsx", "hsk", 4),  # 上下两册合在一个文件里，多一列「册」
]

# 「课文原文」= 只有课文正文、没有挖空的书（HSK5/6 是长篇散文）。每课一篇，
# 只产出 读原文 + 连词成句；vocab/grammar 为 None，练习面板自动隐藏生词挖空。
PROSE_SOURCES = [
    ("HSK5_课文原文.xlsx", "hsk", 5),
    ("HSK6_课文原文.xlsx", "hsk", 6),
]

# 走遍中国：每单元 1 篇听力对话（有说话人）+ 1–3 篇阅读/实用文（书信、通知、明信片…，
# 按 prose 处理）。原文按段落存在「逐段明细」表里（一行一段，按 text_id 归篇），没有挖空。
# 阅读文本原是 OCR，已凭语境人工修正（未对照原书逐字复核）。
DC_SOURCES = [
    ("走遍中国_第四册_课文原文.xlsx", "zoubian", 4),
]

# 生存汉语一本通：每课一段情境对话，按说话轮次存在「对话逐句」表（一行一轮：课号/课名/说话人/中文发言/拼音/英文），
# 没有挖空 → 读原文 + 连词成句。第1–3课是拼音知识课，没有对话。
TURN_SOURCES = [
    ("生存汉语一本通_课文原文.xlsx", "survival", 1),
]

SHEET = "教师答案版"
# Teacher-sheet columns, located by header text (the HSK4 file carries an extra 「册」
# column, so fixed indices would silently read the wrong fields).
COLS = {
    "lesson": ("课次",), "name": ("课名",), "title": ("课文序号",), "orig": ("原文",),
    "vcloze": ("第一篇｜生词/重点词语挖空",), "vans": ("第一篇答案",),
    "gcloze": ("第二篇｜语法挖空", "第二篇｜语法/核心句式挖空"), "gans": ("第二篇答案",),
    "note": ("本课语法点", "本课语言点"),
}

# 中文弯引号要写全角“”‘’——这里曾误写成半角 "" ''（两段字符串被隐式拼接，弯引号根本不在集合里），结果引号被做成了连词成句的方块。
PUNCT = set("，。！？、；：“”‘’「」『』\"'（）()《》…—·～,.!?;: 　\t\n")

VOCAB_BLANK = re.compile(r"[（(]\s*(\d+)\s*[）)]_*")
UNDERSCORES = re.compile(r"_+")


def load_segments() -> dict:
    seg = {}
    if SEG_FILE.exists():
        for line in SEG_FILE.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            clause, toks = line.split("\t", 1)
            seg[clause] = [t for t in toks.split(" ") if t.strip()]
    return seg


def load_gloss() -> dict:
    """Hand-authored pinyin + English per dialogue line (scripts/text_gloss.tsv).
    Keyed by the exact speech text (speaker stripped) — same key the reader shows
    and gen_audio speaks. Lines absent here simply carry no py/en (front-end hides
    the row). Edit the TSV, not the generated JSON."""
    gloss = {}
    if GLOSS_FILE.exists():
        for line in GLOSS_FILE.read_text(encoding="utf-8").splitlines():
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) < 3:
                continue
            zh, py, en = parts[0], parts[1].strip(), parts[2].strip()
            gloss[zh] = {"py": py, "en": en}
    return gloss


def load_word_map() -> dict:
    """hanzi -> {py, en} for the 生词挖空 word-bank tiles. Built from EVERY
    flashcard word (all units), so a cloze answer taught in any lesson shows its
    pinyin + English. scripts/word_gloss.tsv hand-fills the few answers that
    aren't standalone flashcard entries (compounds, single chars, phrases)."""
    m = {}
    for f in DATA.glob("*-*.json"):  # flashcard unit files (texts/ is a subdir)
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue
        if not isinstance(d, dict) or "lessons" not in d:
            continue
        for lesson in d["lessons"]:
            for w in lesson["words"]:
                h = (w.get("hanzi") or "").strip()
                if h and h not in m:
                    m[h] = {"py": w.get("pinyin") or "", "en": w.get("meaning") or ""}
    if WORD_GLOSS_FILE.exists():  # hand overrides/supplements win
        for line in WORD_GLOSS_FILE.read_text(encoding="utf-8").splitlines():
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) < 3:
                continue
            m[parts[0]] = {"py": parts[1].strip(), "en": parts[2].strip()}
    return m


# Real flashcard words that, in the sentences 连词成句 actually tiles, only ever occur as
# the wrong prefix of a longer word (不便宜 → 不便|宜, 从中间 → 从中|间). Excluded from the
# segmentation lexicon ONLY — they stay normal flashcard entries and normal 课文 text.
# Both DO appear standalone elsewhere (「带来不便」「从中了解中国」), but only inside long
# narrative lines that exceed MAX_SENT_TOKENS and are dropped before tiling, so excluding
# them costs nothing today. Before adding a word here, check it against t["sentences"] —
# the tiled items — not just the raw 课文, and re-check when a book is added.
LEX_EXCLUDE = {"不便", "从中", "有机"}  # 有机：中级词表的 organic，课文里只出现在「有机会」中
# 走遍中国4 的词：别的课文里只以「商量好|在」「爱|上中文课」「最|爱看」出现
LEX_EXCLUDE |= {"好在", "爱上", "最爱"}


def load_lexicon(segmap: dict) -> set:
    """Word list for 连词成句 segmentation, built from the app's OWN vocabulary:
    every flashcard word (src/data/{series}-{unit}.json) + every hand-vetted
    grammar-sentence token. Using the taught vocabulary as the lexicon keeps the
    tile boundaries pedagogically aligned and needs no jieba."""
    lex = set()
    for f in DATA.glob("*-*.json"):  # flashcard unit files (texts/ excluded — subdir)
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue
        if not isinstance(d, dict) or "lessons" not in d:
            continue
        for lesson in d["lessons"]:
            for w in lesson["words"]:
                h = (w.get("hanzi") or "").strip()
                if len(h) >= 2:
                    lex.add(h)
    for toks in segmap.values():
        for t in toks:
            if len(t) >= 2:
                lex.add(t)
    return lex - LEX_EXCLUDE


def segment_sentence(text: str, lex: set) -> list:
    """Greedy longest-match word segmentation over `lex`; punctuation is dropped
    (it's not tiled in 连词成句). Unknown runs fall back to single characters."""
    toks = []
    i, n = 0, len(text)
    while i < n:
        ch = text[i]
        if ch in PUNCT:
            i += 1
            continue
        matched = None
        for L in range(min(MAX_WORD, n - i), 1, -1):
            cand = text[i:i + L]
            # a lexicon entry carrying punctuation (会话360 answer 「不行，不行」) would
            # become a tile with a comma in it — never match across punctuation
            if cand in lex and not any(c in PUNCT for c in cand):
                matched = cand
                break
        if matched:
            toks.append(matched)
            i += len(matched)
        else:
            toks.append(ch)
            i += 1
    return toks


def clean(v):
    if v is None:
        return None
    s = str(v).replace("\r\n", "\n").strip()
    return s or None


def parse_vocab(cloze: str, ans: str):
    """Split a numbered-blank cloze into segments + an ordered answer list."""
    if not cloze:
        return None
    seg = []
    pos = 0
    max_n = 0
    for m in VOCAB_BLANK.finditer(cloze):
        if m.start() > pos:
            seg.append({"t": cloze[pos:m.start()]})
        n = int(m.group(1))
        max_n = max(max_n, n)
        seg.append({"b": n})
        pos = m.end()
    if pos < len(cloze):
        seg.append({"t": cloze[pos:]})
    if max_n == 0:
        return None  # no blanks parsed — treat as plain text elsewhere

    answers = [""] * max_n
    for part in re.split(r"[；;\n]", ans or ""):
        part = part.strip()
        mm = re.match(r"^(\d+)\s*[\.、\．]?\s*(.+)$", part)
        if mm:
            idx = int(mm.group(1)) - 1
            if 0 <= idx < max_n:
                answers[idx] = mm.group(2).strip()
    return {"seg": seg, "answers": answers}


def heal_name_blanks(vocab):
    """Never blank the speaker name. The source cloze blanks the FIRST occurrence
    of each answer word, which sometimes lands inside a speaker name (王一[雪]：,
    周[太太]：, [老]同学：). For every blank that falls in a speaker-name prefix
    (before a short ≤8-char colon), un-blank it there; if the same word also occurs
    in the speech body, move the blank to that occurrence, otherwise drop the blank
    entirely (the word is only ever a name — nothing to practise). Blanks are then
    renumbered in reading order. Returns a new vocab dict, or None if no blanks
    survive."""
    seg, answers = vocab["seg"], vocab["answers"]
    atoms = []
    for p in seg:
        if "t" in p:
            atoms += [{"c": ch} for ch in p["t"]]
        else:
            a = answers[p["b"] - 1] if p["b"] - 1 < len(answers) else ""
            atoms.append({"blank": True, "answer": a})

    # Per-line: mark the speaker-name region (before a short colon) and the body.
    lines = [[]]
    for at in atoms:
        (lines.append([]) if at.get("c") == "\n" else lines[-1].append(at))
    for line in lines:
        colon = next((i for i, at in enumerate(line) if at.get("c") in ("：", ":")), None)
        if colon is None:
            for at in line:
                at["inbody"] = True
            continue
        pref = sum(1 if "c" in at else (len(at["answer"] or " ") or 1) for at in line[:colon])
        if pref > 8:  # colon is mid-sentence (narrative), not a speaker delimiter
            for at in line:
                at["inbody"] = True
            continue
        for i, at in enumerate(line):
            at["inbody"] = i > colon
            if i < colon and at.get("blank"):
                at["name"] = True

    body_text = "".join(
        (at["answer"] if at.get("blank") else at["c"])
        for at in atoms
        if at.get("c") != "\n" and at.get("inbody")
    )

    # Un-blank name blanks (restore the word as plain chars); queue relocations.
    out, reloc = [], []
    for at in atoms:
        if at.get("blank") and at.get("name"):
            out += [{"c": ch, "inbody": False} for ch in at["answer"]]
            if at["answer"] and at["answer"] in body_text:
                reloc.append(at["answer"])
        else:
            out.append(at)
    atoms = out

    # Relocate: blank the first body occurrence (a contiguous run of body chars).
    for a in reloc:
        n = len(a)
        pos = [i for i, at in enumerate(atoms) if at.get("inbody") and "c" in at]
        for k in range(len(pos) - n + 1):
            idxs = pos[k:k + n]
            if idxs == list(range(idxs[0], idxs[0] + n)) and "".join(atoms[j]["c"] for j in idxs) == a:
                atoms = atoms[:idxs[0]] + [{"blank": True, "answer": a}] + atoms[idxs[-1] + 1:]
                break

    # Rebuild seg with blanks renumbered in reading order.
    seg_out, ans_out, buf, num = [], [], [], 0
    for at in atoms:
        if at.get("blank"):
            if buf:
                seg_out.append({"t": "".join(buf)})
                buf = []
            num += 1
            seg_out.append({"b": num})
            ans_out.append(at["answer"])
        else:
            buf.append(at["c"])
    if buf:
        seg_out.append({"t": "".join(buf)})
    if num == 0:
        return None
    return {"seg": seg_out, "answers": ans_out}


def strip_punct(tokens):
    return [t for t in tokens if t and t not in PUNCT]


def load_merge_words() -> set:
    """Hand-vetted 补充合并表 (scripts/merge_words.txt): real words / names the lexicon
    lacks, so max-match leaves them as single characters (规|模, 周|瑜)."""
    if not MERGE_FILE.exists():
        return set()
    return {w.strip() for w in MERGE_FILE.read_text(encoding="utf-8").splitlines()
            if w.strip() and not w.startswith("#")}


MERGE_WORDS = load_merge_words()
MERGE_MAX = max((len(w) for w in MERGE_WORDS), default=0)


def merge_tiles(tokens):
    """Post-segmentation clean-up. It only JOINS whole tiles, never splits one, so it
    can't create a tile that straddles a word boundary:
      · adjacent tiles that spell a MERGE_WORDS entry become one (规|模 → 规模);
      · a lone 儿 is erhua — glued onto the tile before it (没事|儿 → 没事儿);
      · consecutive digit tiles form one number (2|0|12 → 2012)."""
    # MERGE_WORDS first, so 儿时 (和|儿|时) is claimed before the erhua rule sees a lone 儿.
    merged, i = [], 0
    while i < len(tokens):
        for k in range(min(len(tokens) - i, MERGE_MAX), 1, -1):
            w = "".join(tokens[i:i + k])
            if w in MERGE_WORDS:
                merged.append(w)
                i += k
                break
        else:
            merged.append(tokens[i])
            i += 1
    out = []
    for t in merged:
        if out and (t == "儿" or (t.isdigit() and out[-1].isdigit())):
            out[-1] += t
        else:
            out.append(t)
    return out


def parse_grammar(cloze: str, ans: str, segmap: dict):
    """One blank spanning a clause → before/after text + answer + word tiles."""
    if not ans:
        return None
    answer = ans.strip()
    tokens = segmap.get(answer)
    if tokens is None:
        # Fallback (should be rare): character-level split, minus punctuation.
        tokens = [c for c in answer]
    tokens = strip_punct(tokens)

    before, after = "", ""
    if cloze:
        runs = list(UNDERSCORES.finditer(cloze))
        if runs:
            before = cloze[: runs[0].start()]
            after = cloze[runs[-1].end():]
    return {"before": before, "after": after, "answer": answer, "tokens": tokens}


SPEAKER_RE = re.compile(r"^([^：:]{1,8})[：:](.*)$")


SENT_SPLIT_RE = re.compile(r"[^。！？；!?;]*[。！？；!?;]+[”\"’'）)」』]*|[^。！？；!?;]+$")


def split_sentences(text: str):
    """Split a speech line into individual sentences by sentence-final punctuation
    (。！？；), keeping the punctuation and any trailing closing quote attached."""
    return [c.strip() for c in SENT_SPLIT_RE.findall(text or "") if c.strip()]


def build_sentences(original: str, segmap: dict, lex: set, answers=None, prose=False):
    """连词成句 items, one SENTENCE each (not one whole line — narratives were a
    single 40–70-tile monster otherwise). Each line is split by sentence-final
    punctuation; a hand-vetted grammar clause (in segmap) is kept whole. Sentences
    that tokenize to <2 tiles (trivial) or >MAX_SENT_TOKENS (too hard) are dropped.
    Sentences containing a 生词挖空 answer word are surfaced FIRST (练完挖空再排句)."""
    answer_set = {a for a in (answers or []) if a}
    primary, rest = [], []
    for raw in (original or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = None if prose else SPEAKER_RE.match(line)
        speaker = m.group(1) if m else ""
        speech = (m.group(2).strip() if m else line)
        if not speech:
            continue
        subs = [speech] if speech in segmap else split_sentences(speech)
        for sub in subs:
            # a hand-vetted segmap entry is final; machine segmentation gets merge_tiles
            toks = strip_punct(segmap[sub]) if sub in segmap else merge_tiles(strip_punct(segment_sentence(sub, lex)))
            if len(toks) < 2 or len(toks) > MAX_SENT_TOKENS:
                continue
            item = {"speaker": speaker, "text": sub, "tokens": toks}
            (primary if any(a in sub for a in answer_set) else rest).append(item)
    return primary + rest


def build_lines(original: str, gloss: dict, prose=False):
    """Per-line reading view: speaker + speech + hand-authored pinyin/English.
    Derived from `original` so the source stays the single dialogue string.
    `prose`: paragraphs of an essay — a leading 「老鼠说：」 is narration, not a speaker."""
    out = []
    for raw in (original or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = None if prose else SPEAKER_RE.match(line)
        name = m.group(1) if m else ""
        speech = (m.group(2).strip() if m else line)
        # 书信/通知的称呼语（「亲爱的同学们：」「家月姐姐：」）长得像说话人但正文为空，
        # 会渲染成一个只有名字的空行 —— 整行当正文处理。
        if not speech:
            name, speech = "", line
        g = gloss.get(speech, {})
        out.append({"name": name, "zh": speech, "py": g.get("py", ""), "en": g.get("en", "")})
    return out


def convert(fname, series, unit, segmap, lex, gloss, wordmap):
    wb = openpyxl.load_workbook(ROOT / "sources" / fname, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "序号")
    header = [clean(h) for h in rows[hi]]
    col = {}
    for k, names in COLS.items():
        col[k] = next((header.index(n) for n in names if n in header), None)
        if col[k] is None:
            raise SystemExit(f"{fname}: 找不到列 {names}")
    C_LESSON, C_NAME, C_TITLE, C_ORIG = col["lesson"], col["name"], col["title"], col["orig"]
    C_VCLOZE, C_VANS, C_GCLOZE, C_GANS, C_NOTE = (col[k] for k in ("vcloze", "vans", "gcloze", "gans", "note"))

    lessons = {}
    order = []
    for r in rows[hi + 1:]:
        if not r or r[C_LESSON] is None:
            continue
        ln = int(r[C_LESSON])
        if ln not in lessons:
            lessons[ln] = {"num": ln, "name": clean(r[C_NAME]) or "", "texts": []}
            order.append(ln)
        vocab = parse_vocab(clean(r[C_VCLOZE]), clean(r[C_VANS]))
        if vocab:  # never blank a speaker name (relocate to body or drop the blank)
            vocab = heal_name_blanks(vocab)
        if vocab:  # bake per-answer pinyin + English for the word-bank tiles
            tiles = {}
            for a in vocab["answers"]:
                a = (a or "").strip()
                if a and a in wordmap and a not in tiles:
                    tiles[a] = wordmap[a]
            if tiles:
                vocab["tiles"] = tiles
        # augment the lexicon with this text's own vocab answers so they tile whole
        seg_lex = lex
        if vocab:
            extra = {a for a in vocab["answers"] if len(a) >= 2}
            if extra:
                seg_lex = lex | extra
        original = clean(r[C_ORIG]) or ""
        lessons[ln]["texts"].append({
            "n": len(lessons[ln]["texts"]) + 1,
            "title": clean(r[C_TITLE]) or f"课文{len(lessons[ln]['texts']) + 1}",
            "original": original,
            "lines": build_lines(original, gloss),
            "vocab": vocab,
            "grammar": parse_grammar(clean(r[C_GCLOZE]), clean(r[C_GANS]), segmap),
            "sentences": build_sentences(original, segmap, seg_lex, vocab["answers"] if vocab else None),
            "note": clean(r[C_NOTE]) or "",
        })

    data = {"series": series, "unit": unit, "lessons": [lessons[n] for n in order]}
    n_texts = sum(len(l["texts"]) for l in data["lessons"])
    return data, len(order), n_texts


def convert_prose(fname, series, unit, segmap, lex, gloss):
    """「课文原文」book: sheet 1, one row per lesson; headers are bilingual and vary
    between files (「lesson_id」 vs 「课号 lesson_id」), so match by substring."""
    wb = openpyxl.load_workbook(ROOT / "sources" / fname, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = [clean(h) or "" for h in rows[0]]

    def find(key):
        i = next((i for i, h in enumerate(header) if key in h), None)
        if i is None:
            raise SystemExit(f"{fname}: 找不到列 {key}")
        return i

    C_ID, C_T, C_BODY = find("lesson_id"), find("title"), find("content")
    lessons = []
    for r in rows[1:]:
        if not r or r[C_ID] is None:
            continue
        # paragraphs are separated by blank lines; one paragraph = one reading line
        original = "\n".join(p.strip() for p in (clean(r[C_BODY]) or "").split("\n") if p.strip())
        lessons.append({"num": int(r[C_ID]), "name": clean(r[C_T]) or "", "texts": [{
            "n": 1,
            "title": "课文",
            "original": original,
            "lines": build_lines(original, gloss, prose=True),
            "vocab": None,
            "grammar": None,
            "sentences": build_sentences(original, segmap, lex, prose=True),
            "note": "",
        }]})
    return {"series": series, "unit": unit, "lessons": lessons}, len(lessons), len(lessons)


def convert_dc(fname, series, unit, segmap, lex, gloss):
    """走遍中国 book: 「课文主表」 = one row per text (unit, lesson, material_type, title),
    「逐段明细」 = one row per paragraph, keyed by text_id. A unit's texts are its
    Lesson 1 dialogue followed by its Lesson 2 readings, in sheet order."""
    wb = openpyxl.load_workbook(ROOT / "sources" / fname, read_only=True, data_only=True)

    def table(name):
        rows = list(wb[name].iter_rows(values_only=True))
        header = [clean(h) or "" for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:] if r and r[0]]

    # 拼音/英文 live in the 逐段明细 sheet itself (speech only, speaker stripped) — this
    # book's gloss is NOT in text_gloss.tsv. A blank cell falls back to the TSV.
    paras = {}
    for p in table("逐段明细"):
        zh = clean(p["paragraph_text"])
        if zh:
            paras.setdefault(p["text_id"], []).append((zh, clean(p.get("拼音")), clean(p.get("英文"))))
    # speaker names (永民, 史蒂夫…) join the lexicon so they tile whole, as in convert_360
    names = {m.group(1) for ps in paras.values() for zh, _, _ in ps if (m := SPEAKER_RE.match(zh))}
    lex = lex | {n for n in names if len(n) >= 2}
    lessons = {}
    for t in table("课文主表"):
        ln = int(t["unit"])
        L = lessons.setdefault(ln, {"num": ln, "name": clean(t["unit_title"]) or "", "texts": []})
        prose = t["material_type"] != "听力对话"
        ps = paras.get(t["text_id"], [])
        original = "\n".join(zh for zh, _, _ in ps)
        lines = build_lines(original, gloss, prose=prose)
        for line, (_, py, en) in zip(lines, ps):  # build_lines keeps one line per paragraph
            line["py"], line["en"] = py or line["py"], en or line["en"]
        L["texts"].append({
            "n": len(L["texts"]) + 1,
            "title": clean(t["text_title"]) or f"课文{len(L['texts']) + 1}",
            "original": original,
            "lines": lines,
            "vocab": None,
            "grammar": None,
            "sentences": build_sentences(original, segmap, lex, prose=prose),
            "note": "",
        })
    n_texts = sum(len(L["texts"]) for L in lessons.values())
    return {"series": series, "unit": unit, "lessons": [lessons[n] for n in sorted(lessons)]}, len(lessons), n_texts


def convert_turns(fname, series, unit, segmap, lex, gloss):
    """「对话逐句」book (生存汉语): one row per speaking turn, grouped by 课号 into ONE text
    per lesson. 拼音/英文 come from the sheet's own columns (TSV fallback)."""
    wb = openpyxl.load_workbook(ROOT / "sources" / fname, read_only=True, data_only=True)
    rows = list(wb["对话逐句"].iter_rows(values_only=True))
    header = [clean(h) or "" for h in rows[0]]
    turns = [dict(zip(header, r)) for r in rows[1:] if r and r[0]]
    names = {clean(t["说话人"]) for t in turns if clean(t["说话人"])}
    lex = lex | {n for n in names if len(n) >= 2}  # 杨无瑕/服务员 tile whole
    lessons = {}
    for t in turns:
        zh = clean(t["中文发言"])
        if not zh:
            continue
        ln = int(t["课号"])
        L = lessons.setdefault(ln, {"num": ln, "name": clean(t["课名"]) or "", "turns": []})
        L["turns"].append((clean(t["说话人"]), zh, clean(t.get("拼音")), clean(t.get("英文"))))
    out = []
    for ln in sorted(lessons):
        L = lessons[ln]
        original = "\n".join(f"{who}：{zh}" if who else zh for who, zh, _, _ in L["turns"])
        lines = build_lines(original, gloss)
        for line, (_, _, py, en) in zip(lines, L["turns"]):
            line["py"], line["en"] = py or line["py"], en or line["en"]
        out.append({"num": ln, "name": L["name"], "texts": [{
            "n": 1, "title": "情境对话", "original": original, "lines": lines,
            "vocab": None, "grammar": None,
            "sentences": build_sentences(original, segmap, lex), "note": "",
        }]})
    return {"series": series, "unit": unit, "lessons": out}, len(out), len(out)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    segmap = load_segments()
    lex = load_lexicon(segmap)
    gloss = load_gloss()
    wordmap = load_word_map()
    meta = {}
    print(f"Segmentation table: {len(segmap)} clauses · lexicon: {len(lex)} words · gloss: {len(gloss)} lines · wordmap: {len(wordmap)} words")
    print("Converted 课文:")
    # Several books may map to one unit (第四册 = 上 + 下): accumulate before writing,
    # else the second file would silently overwrite the first.
    units = {}
    jobs = [(f, s, u, lambda f, s, u: convert(f, s, u, segmap, lex, gloss, wordmap)) for f, s, u in SOURCES]
    jobs += [(f, s, u, lambda f, s, u: convert_prose(f, s, u, segmap, lex, gloss)) for f, s, u in PROSE_SOURCES]
    jobs += [(f, s, u, lambda f, s, u: convert_dc(f, s, u, segmap, lex, gloss)) for f, s, u in DC_SOURCES]
    jobs += [(f, s, u, lambda f, s, u: convert_turns(f, s, u, segmap, lex, gloss)) for f, s, u in TURN_SOURCES]
    for fname, series, unit, run in jobs:
        data, n_lessons, n_texts = run(fname, series, unit)
        key = f"{series}-{unit}"
        if key in units:
            units[key]["data"]["lessons"].extend(data["lessons"])
            units[key]["texts"] += n_texts
            units[key]["from"].append(fname)
        else:
            units[key] = {"data": data, "texts": n_texts, "from": [fname]}

    for key, u in units.items():
        data = u["data"]
        data["lessons"].sort(key=lambda l: l["num"])
        (OUT / f"{key}.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        meta[key] = {"lessonCount": len(data["lessons"]), "textCount": u["texts"]}
        print(f"  {key}: {len(data['lessons'])} lessons, {u['texts']} texts  ← {' + '.join(u['from'])}")

    (OUT / "meta.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(f"  meta → {(OUT / 'meta.json').relative_to(ROOT)} ({len(meta)} units)")


if __name__ == "__main__":
    main()
