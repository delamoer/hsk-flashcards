#!/usr/bin/env python3
"""Convert the 会话360 课文挖空练习 xlsx files into per-unit 课文 JSON.

Run with uv (openpyxl + pypinyin provided on the fly):
    uv run --with openpyxl --with pypinyin python3 scripts/convert_360.py

拼音 for 句3/4 (no pinyin column in the source) is generated with pypinyin
(tone-marked); 句1/2 keep their human pinyin. English is hand/AI-authored in
scripts/text_gloss_360.tsv (汉字<TAB>English, keyed by the exact speech line) and
merged into each line's `en` — blank where a line isn't glossed yet.

Sources (sources/汉语标准会话360句{1,2,3,4}_课文挖空练习整理_第1-8课*.xlsx):
  four books → series "huihua360", units 1..4 (matching the flashcard courses,
  课次 1..8 == our lesson num).

Unlike the 逐篇课文双重挖空 files (convert_texts.py), the 360 files have a
heterogeneous schema (different sheet names, some with pinyin, some with
{{blank}} website templates, 句4 has no 练习类型 column at all). So instead of
trusting the pre-blanked cloze text, we REBUILD both practice layers from the
clean dialogue + the answer list, mirroring the existing 套路:
  · 生词挖空 (vocab)  — blank the WORD/PHRASE answers over the full dialogue,
                        fill from a word bank (exactly like convert_texts).
  · 连词成句 (sentences) — split every dialogue line into sentences and tokenize
                        via the shared lexicon (build_sentences, reused).

Word- vs sentence-level answers are separated by average answer length per
exercise (≤ WORD_AVG chars → word-level → feeds vocab; longer → a whole-clause
"整句挖空", skipped for vocab since 连词成句 already drills it).

Output (merged into the same folder / meta as convert_texts, HSK entries kept):
  src/data/texts/huihua360-{1,2,3,4}.json
  src/data/texts/meta.json   (huihua360-* keys merged in)
"""
import json
import re
import statistics
from pathlib import Path

import openpyxl
from pypinyin import Style, pinyin as _pinyin

import convert_texts as ct  # reuse lexicon / wordmap / build_sentences / PUNCT / SPEAKER_RE

GLOSS_TSV = Path(__file__).resolve().parent / "text_gloss_360.tsv"

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "data" / "texts"
DATA = ROOT / "src" / "data"

# (filename, unit) — series is always huihua360.
SOURCES = [
    ("汉语标准会话360句1_课文挖空练习整理_第1-8课_拼音版.xlsx", 1),
    ("汉语标准会话360句2_课文挖空练习整理_第1-8课_拼音版.xlsx", 2),
    ("汉语标准会话360句3_课文挖空练习整理_第1-8课.xlsx", 3),
    ("汉语标准会话360句4_课文挖空练习整理_第1-8课.xlsx", 4),
]
SERIES = "huihua360"

WORD_AVG = 4.0    # exercise counts as word-level if mean(answer len) ≤ this
MAX_BLANK_LEN = 7  # never blank an answer longer than this in the vocab cloze
ANS_SPLIT = re.compile(r"[；;]")


def col(hdr, *names):
    for i, h in enumerate(hdr):
        if h and any(n in str(h) for n in names):
            return i
    return None


def load_lesson_names(unit):
    """Lesson (课次) names from the flashcard unit, so 课文 lesson headers match."""
    try:
        d = json.loads((DATA / f"{SERIES}-{unit}.json").read_text(encoding="utf-8"))
    except Exception:
        return {}
    return {l["num"]: (l.get("name") or "") for l in d.get("lessons", [])}


def read_dialogues(wb):
    """(课次,会话) -> {'zh': dialogue, 'py': pinyin|'', 'title': 课题/场景|''}."""
    ws = wb["对话原文"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    c_les = col(hdr, "课次")
    c_conv = col(hdr, "会话")
    c_zh = col(hdr, "完整对话（汉字）", "完整对话(汉字)")
    c_py = col(hdr, "完整对话（拼音）", "完整对话(拼音)")
    c_title = col(hdr, "课题")
    c_scene = col(hdr, "场景")
    out = {}
    for r in rows[1:]:
        if not r or r[c_les] is None or r[c_conv] is None:
            continue
        zh = ct.clean(r[c_zh]) if c_zh is not None else None
        if not zh:
            continue
        out[(int(r[c_les]), int(r[c_conv]))] = {
            "zh": zh,
            "py": (ct.clean(r[c_py]) if c_py is not None else None) or "",
            "title": (ct.clean(r[c_title]) if c_title is not None else None)
                     or (ct.clean(r[c_scene]) if c_scene is not None else None) or "",
        }
    return out


def find_exercise_sheet(wb):
    for name in wb.sheetnames:
        ws = wb[name]
        first = next(ws.iter_rows(values_only=True), None)
        if first and col(first, "参考答案") is not None:
            return ws
    return None


def read_word_answers(wb):
    """(课次,会话) -> ordered de-duped list of WORD/PHRASE answers to blank.
    Only exercises whose mean answer length ≤ WORD_AVG contribute (word-level);
    clause-level 整句挖空 exercises are left to 连词成句."""
    ws = find_exercise_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    c_les = col(hdr, "课次")
    c_conv = col(hdr, "会话")
    c_ans = col(hdr, "参考答案")
    out = {}
    for r in rows[1:]:
        if not r or r[c_les] is None or r[c_conv] is None or r[c_ans] is None:
            continue
        ans = [a.strip() for a in ANS_SPLIT.split(str(r[c_ans])) if a.strip()]
        if not ans:
            continue
        if statistics.mean(len(a) for a in ans) > WORD_AVG:
            continue  # sentence-level exercise → skip for vocab
        key = (int(r[c_les]), int(r[c_conv]))
        bucket = out.setdefault(key, [])
        for a in ans:
            if 1 <= len(a) <= MAX_BLANK_LEN and a not in bucket:
                bucket.append(a)
    return out


def build_vocab(dialogue: str, answers: list):
    """Blank the FIRST speech-body occurrence of each answer over the full
    dialogue → {seg:[{t}/{b}], answers:[]}. Speaker names (before a short colon)
    are never blanked. Returns None if nothing could be blanked."""
    if not answers:
        return None
    remaining = list(answers)
    seg, ans_out, buf, num = [], [], "", 0

    def flush():
        nonlocal buf
        if buf:
            seg.append({"t": buf})
            buf = ""

    for li, raw in enumerate(dialogue.split("\n")):
        line = raw.rstrip()
        if li > 0:
            buf += "\n"
        m = ct.SPEAKER_RE.match(line)
        if m and len(m.group(1)) <= 8:
            sep = line.find("：")
            if sep < 0:
                sep = line.find(":")
            buf += line[: sep + 1]
            speech = line[sep + 1:]
        else:
            speech = line
        i = 0
        while i < len(speech):
            hit = None
            for a in sorted(remaining, key=len, reverse=True):
                if speech.startswith(a, i):
                    hit = a
                    break
            if hit:
                flush()
                num += 1
                seg.append({"b": num})
                ans_out.append(hit)
                remaining.remove(hit)
                i += len(hit)
            else:
                buf += speech[i]
                i += 1
    flush()
    if num == 0:
        return None
    return {"seg": seg, "answers": ans_out}


def strip_speaker_py(py_raw: str) -> str:
    """Strip a leading 'Name:' from a pinyin line. Romanized names can exceed the
    zh SPEAKER_RE 8-char limit (e.g. 'Tóngxuémen'), so split on the first colon."""
    i = py_raw.find("：")
    if i < 0:
        i = py_raw.find(":")
    return py_raw[i + 1:].strip() if 0 <= i <= 24 else py_raw


def to_pinyin(zh: str) -> str:
    """Tone-marked pinyin for a Chinese line (pypinyin), syllables space-joined,
    punctuation attached without a leading space — matches the human 句1/2 format."""
    out = ""
    for grp in _pinyin(zh, style=Style.TONE, errors=lambda s: list(s)):
        s = grp[0]
        if re.search(r"[A-Za-zɡü]", s):  # a syllable (or latin/digit) → space-separated
            out += ((" " if out and not out.endswith(" ") else "") + s)
        else:  # punctuation / untranslatable → attach
            out += s
    return out.strip()


def build_lines(dialogue: str, pinyin: str, gloss: dict):
    """Per-line reading view {name, zh, py, en}. Pinyin: source line (句1/2) when
    present, else pypinyin (句3/4). English: from the hand/AI gloss (blank if未收录)."""
    zh_lines = [l.strip() for l in dialogue.split("\n") if l.strip()]
    py_lines = [l.strip() for l in (pinyin or "").split("\n") if l.strip()]
    out = []
    for i, line in enumerate(zh_lines):
        m = ct.SPEAKER_RE.match(line)
        name = m.group(1) if m else ""
        zh = (m.group(2).strip() if m else line)
        py_raw = py_lines[i] if i < len(py_lines) else ""
        py = strip_speaker_py(py_raw) or to_pinyin(zh)
        out.append({"name": name, "zh": zh, "py": py, "en": gloss.get(zh, "")})
    return out


def load_gloss() -> dict:
    d = {}
    if GLOSS_TSV.exists():
        for line in GLOSS_TSV.read_text(encoding="utf-8").splitlines():
            if not line.strip() or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) >= 2 and parts[0].strip() and parts[1].strip():
                d[parts[0].strip()] = parts[1].strip()
    return d


def dialogue_names(dialogues) -> set:
    """Speaker names across a book's dialogues — added to the 连词成句 lexicon so a
    name tiles whole (a name-only sentence then has 1 token and is dropped, instead
    of fragmenting into single-char 排序 items like 雪+丽)."""
    names = set()
    for d in dialogues.values():
        for raw in d["zh"].split("\n"):
            m = ct.SPEAKER_RE.match(raw.strip())
            if m and len(m.group(1)) >= 2:
                names.add(m.group(1))
    return names


def convert(fname, unit, segmap, lex, wordmap, gloss):
    wb = openpyxl.load_workbook(ROOT / "sources" / fname, read_only=True, data_only=True)
    dialogues = read_dialogues(wb)
    wanswers = read_word_answers(wb)
    names = load_lesson_names(unit)
    name_lex = lex | dialogue_names(dialogues)

    lessons = {}
    order = []
    for (les, conv) in sorted(dialogues):
        d = dialogues[(les, conv)]
        if les not in lessons:
            lessons[les] = {"num": les, "name": names.get(les, ""), "texts": []}
            order.append(les)
        vocab = build_vocab(d["zh"], wanswers.get((les, conv), []))
        if vocab:
            tiles = {}
            for a in vocab["answers"]:
                if a in wordmap and a not in tiles:
                    tiles[a] = wordmap[a]
            if tiles:
                vocab["tiles"] = tiles
        seg_lex = name_lex
        if vocab:
            extra = {a for a in vocab["answers"] if len(a) >= 2}
            if extra:
                seg_lex = name_lex | extra
        n = len(lessons[les]["texts"]) + 1
        title = d["title"] or f"会话 {conv}"
        lessons[les]["texts"].append({
            "n": n,
            "title": title,
            "original": d["zh"],
            "lines": build_lines(d["zh"], d["py"], gloss),
            "vocab": vocab,
            "grammar": None,
            "sentences": ct.build_sentences(d["zh"], segmap, seg_lex,
                                            vocab["answers"] if vocab else None),
            "note": "",
        })

    data = {"series": SERIES, "unit": unit, "lessons": [lessons[n] for n in order]}
    n_texts = sum(len(l["texts"]) for l in data["lessons"])
    return data, len(order), n_texts


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    segmap = ct.load_segments()
    lex = ct.load_lexicon(segmap)
    wordmap = ct.load_word_map()
    gloss = load_gloss()

    meta_path = OUT / "meta.json"
    meta = json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}

    print(f"lexicon: {len(lex)} words · wordmap: {len(wordmap)} words · gloss(EN): {len(gloss)} lines")
    print("Converted 会话360 课文:")
    for fname, unit in SOURCES:
        data, n_lessons, n_texts = convert(fname, unit, segmap, lex, wordmap, gloss)
        key = f"{SERIES}-{unit}"
        (OUT / f"{key}.json").write_text(
            json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8"
        )
        meta[key] = {"lessonCount": n_lessons, "textCount": n_texts}
        print(f"  {key}: {n_lessons} lessons, {n_texts} texts  ← {fname}")

    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"  meta → {meta_path.relative_to(ROOT)} ({len(meta)} units)")


if __name__ == "__main__":
    main()
